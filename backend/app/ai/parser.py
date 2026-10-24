import os
import re
import hashlib
from typing import Tuple, Optional, Dict
from datetime import datetime
import io

def compute_checksum(file_bytes: bytes) -> str:
    return hashlib.sha256(file_bytes).hexdigest()

def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_content = []
    
    # Try pdfplumber first
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                # Try layout mode for multi-column resumes, fallback to standard flow
                page_text = page.extract_text(layout=True)
                if not page_text or len(page_text.strip()) < 10:
                    page_text = page.extract_text()
                if page_text:
                    text_content.append(page_text)
        if text_content:
            return "\n".join(text_content).strip()
    except Exception:
        pass

    # Fallback to pypdf
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_content.append(page_text)
        if text_content:
            return "\n".join(text_content).strip()
    except Exception:
        pass

    return ""

def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs).strip()
    except Exception:
        return ""

def parse_document(file_name: str, file_bytes: bytes) -> Tuple[str, str, int, str]:
    """
    Parses document bytes based on file extension.
    Returns (extracted_text, file_type, file_size_kb, checksum)
    """
    file_size_kb = max(1, len(file_bytes) // 1024)
    checksum = compute_checksum(file_bytes)
    ext = os.path.splitext(file_name)[1].lower()

    extracted_text = ""
    file_type = "txt"

    if ext == ".pdf":
        file_type = "pdf"
        extracted_text = extract_text_from_pdf(file_bytes)
    elif ext in [".docx", ".doc"]:
        file_type = "docx"
        extracted_text = extract_text_from_docx(file_bytes)
    else:
        file_type = "txt"
        try:
            extracted_text = file_bytes.decode("utf-8", errors="ignore")
        except Exception:
            extracted_text = str(file_bytes)

    if not extracted_text.strip():
        # Fallback decode
        try:
            extracted_text = file_bytes.decode("latin-1", errors="ignore")
        except Exception:
            extracted_text = "Empty or unparseable document content."

    return extracted_text.strip(), file_type, file_size_kb, checksum


_UNIVERSITY_RE = re.compile(
    r"(?:[A-Z][A-Za-z&.'-]+[ \t]+){0,5}(?:University|College|Institute)"
    r"(?:[ \t]+of[ \t]+[A-Z][A-Za-z&.'-]+(?:[ \t]+(?:and[ \t]+)?[A-Z][A-Za-z&.'-]+){0,4})?"
    r"|جامعة[ \t]+[\u0600-\u06FF]+(?:[ \t]+[\u0600-\u06FF]+)?"
)
_MAJOR_DEGREE_RE = re.compile(
    r"\b(?:Bachelor|Master|B\.?\s?Sc|M\.?\s?Sc|Diploma|Degree)\b[^\n|]*?\bin\s+"
    r"([A-Z][A-Za-z&/ ]+?)\s*(?=\||,|\(|-|–|\bat\b|$)",
    re.MULTILINE,
)
_MAJOR_LABEL_RE = re.compile(r"\b(?:Major|Specialization)\s*[:\-]\s*([^\n|,]+)", re.IGNORECASE)
_GPA_RE = re.compile(
    r"\bGPA\b\s*[:\-]?\s*(\d(?:\.\d{1,2})?)(?:\s*(?:/|out of)\s*(\d(?:\.\d{1,2})?))?",
    re.IGNORECASE,
)
_CAREER_LEVELS = [
    (re.compile(r"\b(?:fresh\s+graduate|خريج\s+جديد|حديث\s+التخرج)\b", re.IGNORECASE), "Fresh Graduate"),
    (re.compile(r"\bsenior\s+(?:\w+\s+){0,3}student\b|\bطالب\s+سنة\s+أخيرة\b", re.IGNORECASE), "Senior Student"),
    (re.compile(r"\b(?:junior|undergraduate)\s+(?:\w+\s+){0,3}student\b|\bexpected\s+graduation\b|\bطالب\s+جامعي\b", re.IGNORECASE), "Student"),
    (re.compile(r"\b(?:graduated|alumn(?:us|a|i))\b", re.IGNORECASE), "Graduate"),
]
_GRAD_DATE_RE = re.compile(
    r"\b(?:graduation|expected\s+graduation|تاريخ\s+التخرج)\s*[:\-]?\s*(?:[A-Za-z]+\s+)?(\d{4})\b",
    re.IGNORECASE
)
_STUDENT_ID_RE = re.compile(
    r"\b(?:Student|University)[ \t]*(?:ID|No\.?|Number)[ \t]*[:#\-]?[ \t]*([A-Za-z0-9-]{4,20})",
    re.IGNORECASE,
)
_SUMMARY_HEADING_RE = re.compile(
    r"^[ \t]*(?:PROFESSIONAL[ \t]+SUMMARY|CAREER[ \t]+OBJECTIVE|SUMMARY|PROFILE|OBJECTIVE|ABOUT[ \t]+ME|الملخص[ \t]+المهني|ملخص[ \t]+مهني|نبذة[ \t]+شخصية)[ \t]*:?[ \t]*\r?$",
    re.IGNORECASE | re.MULTILINE,
)


def _extract_summary(text: str) -> Optional[str]:
    """Returns the paragraph under a Summary/Profile/Objective heading, up to a blank line or the next heading."""
    heading = _SUMMARY_HEADING_RE.search(text)
    if not heading:
        return None
    lines = []
    for line in text[heading.end():].lstrip("\r\n").splitlines():
        stripped = line.strip()
        if not stripped or (stripped.isupper() and len(stripped) > 3):
            break
        lines.append(stripped)
    summary = " ".join(lines)
    return summary[:1000] or None


def extract_profile_fields(text: str) -> Dict[str, Optional[str]]:
    """
    Extracts candidate profile fields (university, major, gpa, career_level,
    student_id, profile_summary) from resume text using regex heuristics.
    Any field that cannot be found is None.
    """
    text = text or ""

    uni_match = _UNIVERSITY_RE.search(text)
    university = uni_match.group(0).strip() if uni_match else None

    major = None
    major_match = _MAJOR_DEGREE_RE.search(text) or _MAJOR_LABEL_RE.search(text)
    if major_match:
        major = major_match.group(1).strip()[:255] or None

    gpa = None
    gpa_match = _GPA_RE.search(text)
    if gpa_match:
        gpa = f"{gpa_match.group(1)} / {gpa_match.group(2)}" if gpa_match.group(2) else gpa_match.group(1)

    career_level = next((level for pattern, level in _CAREER_LEVELS if pattern.search(text)), None)
    if not career_level:
        grad_match = _GRAD_DATE_RE.search(text)
        if grad_match:
            try:
                grad_year = int(grad_match.group(1))
                curr_year = datetime.now().year
                if grad_year > curr_year:
                    career_level = "Senior Student"
                elif grad_year >= curr_year - 1:
                    career_level = "Fresh Graduate"
                else:
                    career_level = "Graduate"
            except Exception:
                pass
    if not career_level and re.search(r"\b(?:graduate|alumn(?:us|a|i)|خريج)\b", text, re.IGNORECASE):
        career_level = "Fresh Graduate"

    student_id_match = _STUDENT_ID_RE.search(text)

    return {
        "university": university[:255] if university else None,
        "major": major,
        "gpa": gpa,
        "career_level": career_level,
        "student_id": student_id_match.group(1) if student_id_match else None,
        "profile_summary": _extract_summary(text),
    }
